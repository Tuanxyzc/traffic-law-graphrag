"""Unit tests for the Deterministic Benchmark Engine (Spec-Driven Development).

Validates:
1. Pydantic schemas (DeterministicBenchmarkItem, SanctionSlot, etc.) and legacy adapters.
2. Mode 1: Retrieval Set-Theoretic evaluation (must-have recall, forbidden collision, optional overlap).
3. Mode 2: Generation Fact-checking (slot-filling, regex matching, negative noise filtering, behavioral flags).
4. Failure Root-Cause Attribution classification.
5. EvaluationOrchestrator aggregation in both modes.
6. Excel & JSON report exporter.
"""

from __future__ import annotations

from pathlib import Path
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook

from src.eval.deterministic import (
    evaluate_benchmark_item,
    evaluate_generation,
    evaluate_retrieval,
    extract_license_suspension,
    extract_points_deducted,
)
from src.eval.evaluator import EvaluationOrchestrator
from src.eval.exporter import EvaluationExporter
from src.eval.models import (
    BenchmarkItemEvaluationResult,
    BenchmarkRunSummary,
    DeterministicBenchmarkItem,
    FactualGroundTruth,
    FailureRootCause,
    GenerationAssertions,
    GenerationEvaluationResult,
    RetrievalEvaluationResult,
    RetrievalGroundTruth,
    SanctionSlot,
    TestCategory,
)

# ============================================================================
# Fixtures
# ============================================================================


@pytest.fixture
def sample_fine_lookup_item() -> DeterministicBenchmarkItem:
    return DeterministicBenchmarkItem(
        test_id="TC_TEST_001",
        category=TestCategory.FINE_LOOKUP,
        question="Người đi xe máy vượt đèn đỏ bị phạt bao nhiêu tiền và trừ mấy điểm?",
        retrieval_gt=RetrievalGroundTruth(
            must_have_node_ids=["168_2024_ND-CP_D7_K7_Dc"],
            optional_node_ids=["36_2024_QH15_D11_K4"],
            must_not_have_node_ids=["168_2024_ND-CP_D6_K9_Db"],
        ),
        factual_gt=FactualGroundTruth(
            vehicle_sanctions={
                "xe_mo_to": SanctionSlot(
                    fine_min=4_000_000,
                    fine_max=6_000_000,
                    points_deducted=4,
                    license_suspended_months_min=None,
                    license_suspended_months_max=None,
                )
            },
            required_citations=["Điểm c Khoản 7 Điều 7", "Nghị định 168/2024/NĐ-CP"],
            exact_dates=[],
        ),
        assertions=GenerationAssertions(
            must_contain_regex=[
                r"(4[\.,]000[\.,]000|4\s*triệu)",
                r"(6[\.,]000[\.,]000|6\s*triệu)",
                r"trừ\s*(4|bốn)\s*điểm",
            ],
            must_not_contain_keywords=["xe ô tô", "18.000.000", "20.000.000"],
            is_refusal_expected=False,
            is_amended_warning_expected=False,
        ),
        notes="Test case kiểm tra vượt đèn đỏ xe máy",
    )


@pytest.fixture
def sample_out_of_scope_item() -> DeterministicBenchmarkItem:
    return DeterministicBenchmarkItem(
        test_id="TC_TEST_OUT_SCOPE",
        category=TestCategory.OUT_OF_SCOPE,
        question="Quy định xử phạt vượt đèn đỏ đối với tàu hỏa đường sắt như thế nào?",
        retrieval_gt=RetrievalGroundTruth(
            must_have_node_ids=[],
            optional_node_ids=[],
            must_not_have_node_ids=["168_2024_ND-CP_D6_K9_Db"],
        ),
        factual_gt=FactualGroundTruth(),
        assertions=GenerationAssertions(
            must_contain_regex=[],
            must_not_contain_keywords=[],
            is_refusal_expected=True,
            is_amended_warning_expected=False,
        ),
        notes="Câu hỏi ngoài phạm vi đường bộ",
    )


@pytest.fixture
def sample_amended_item() -> DeterministicBenchmarkItem:
    return DeterministicBenchmarkItem(
        test_id="TC_TEST_AMENDED",
        category=TestCategory.DOCUMENT_AMENDMENT,
        question="Nghị định 100/2019/NĐ-CP còn hiệu lực áp dụng xử phạt giao thông không?",
        retrieval_gt=RetrievalGroundTruth(
            must_have_node_ids=["168_2024_ND-CP_D52_K1"],
            optional_node_ids=[],
            must_not_have_node_ids=[],
        ),
        factual_gt=FactualGroundTruth(
            required_citations=["Nghị định 168/2024/NĐ-CP"],
        ),
        assertions=GenerationAssertions(
            must_contain_regex=[r"(hết\s+hiệu\s+lực|bãi\s+bỏ|thay\s+thế)"],
            must_not_contain_keywords=[],
            is_refusal_expected=False,
            is_amended_warning_expected=True,
        ),
        notes="Test case kiểm tra cảnh báo hiệu lực văn bản đã hết",
    )


# ============================================================================
# 1. Pydantic Models & Backward Compatibility Tests
# ============================================================================


def test_sanction_slot_validation() -> None:
    slot = SanctionSlot(
        fine_min=8_000_000,
        fine_max=10_000_000,
        points_deducted=12,
        license_suspended_months_min=10,
        license_suspended_months_max=12,
    )
    assert slot.fine_min == 8_000_000
    assert slot.fine_max == 10_000_000
    assert slot.points_deducted == 12
    assert slot.license_suspended_months_min == 10
    assert slot.license_suspended_months_max == 12


def test_deterministic_benchmark_item_to_legacy(
    sample_fine_lookup_item: DeterministicBenchmarkItem,
) -> None:
    legacy = sample_fine_lookup_item.to_legacy_sample()
    assert legacy.id == "TC_TEST_001"
    assert legacy.category.value == "FINE_LOOKUP"
    assert legacy.expected_fine_min == 4_000_000
    assert legacy.expected_fine_max == 6_000_000
    assert legacy.expected_points_deducted == 4
    assert legacy.expected_provision_ids == ["168_2024_ND-CP_D7_K7_Dc"]


# ============================================================================
# 2. Mode 1: Retrieval Set-Theoretic Evaluation Tests
# ============================================================================


def test_evaluate_retrieval_perfect_match(
    sample_fine_lookup_item: DeterministicBenchmarkItem,
) -> None:
    retrieved = [
        "168_2024_ND-CP_D7_K7_Dc",
        "36_2024_QH15_D11_K4",
        "168_2024_ND-CP_D7_K1",
    ]
    res: RetrievalEvaluationResult = evaluate_retrieval(
        sample_fine_lookup_item.retrieval_gt, retrieved
    )
    assert res.passed is True
    assert res.recall == 1.0
    assert res.collision_rate == 0.0
    assert res.must_have_matched == ["168_2024_ND-CP_D7_K7_Dc"]
    assert res.must_have_missing == []
    assert res.must_not_have_detected == []
    assert "36_2024_QH15_D11_K4" in res.optional_matched


def test_evaluate_retrieval_missing_must_have(
    sample_fine_lookup_item: DeterministicBenchmarkItem,
) -> None:
    retrieved = ["36_2024_QH15_D11_K4", "168_2024_ND-CP_D7_K1"]
    res = evaluate_retrieval(sample_fine_lookup_item.retrieval_gt, retrieved)
    assert res.passed is False
    assert res.recall == 0.0
    assert res.must_have_missing == ["168_2024_ND-CP_D7_K7_Dc"]
    assert res.collision_rate == 0.0


def test_evaluate_retrieval_forbidden_collision(
    sample_fine_lookup_item: DeterministicBenchmarkItem,
) -> None:
    # Contains both must-have and forbidden node (car violation)
    retrieved = ["168_2024_ND-CP_D7_K7_Dc", "168_2024_ND-CP_D6_K9_Db"]
    res = evaluate_retrieval(sample_fine_lookup_item.retrieval_gt, retrieved)
    assert res.passed is False
    assert res.recall == 1.0
    assert res.collision_rate == 1.0
    assert res.must_not_have_detected == ["168_2024_ND-CP_D6_K9_Db"]


# ============================================================================
# 3. Mode 2: Generation Evaluation & Fact-Checking Tests
# ============================================================================


def test_extract_points_deducted() -> None:
    text1 = "Người điều khiển xe bị phạt tiền và bị trừ 4 điểm trên giấy phép lái xe."
    assert extract_points_deducted(text1) == 4

    text2 = "Hành vi vi phạm này sẽ bị trừ mười hai điểm GPLX."
    assert extract_points_deducted(text2) == 12

    text3 = "Hành vi này chỉ bị phạt tiền, không bị trừ điểm GPLX."
    assert extract_points_deducted(text3) == 0

    text4 = "Phạt tiền từ 1 triệu đến 2 triệu đồng."
    assert extract_points_deducted(text4) is None


def test_extract_license_suspension() -> None:
    text1 = "Bị tước quyền sử dụng Giấy phép lái xe từ 10 tháng đến 12 tháng."
    assert extract_license_suspension(text1) == (10, 12)

    text2 = "Bị tước bằng lái xe 2 tháng kể từ ngày ra quyết định."
    assert extract_license_suspension(text2) == (2, 2)

    text3 = "Không áp dụng hình thức tước quyền sử dụng GPLX."
    assert extract_license_suspension(text3) == (None, None)


def test_evaluate_generation_passed(
    sample_fine_lookup_item: DeterministicBenchmarkItem,
) -> None:
    answer = (
        "Theo Điểm c Khoản 7 Điều 7 Nghị định 168/2024/NĐ-CP, người điều khiển xe mô tô, xe gắn máy "
        "vượt đèn đỏ bị phạt tiền từ 4.000.000 đồng đến 6.000.000 đồng và bị trừ 4 điểm giấy phép lái xe."
    )
    res: GenerationEvaluationResult = evaluate_generation(
        ground_truth=sample_fine_lookup_item.factual_gt,
        assertions=sample_fine_lookup_item.assertions,
        generated_answer=answer,
    )
    assert res.passed is True
    assert res.vehicle_sanctions_passed.get("xe_mo_to") is True
    assert all(res.regex_passed.values())
    assert len(res.forbidden_keywords_detected) == 0
    assert res.citations_passed is True
    assert res.refusal_passed is True


def test_evaluate_generation_forbidden_keyword_leak(
    sample_fine_lookup_item: DeterministicBenchmarkItem,
) -> None:
    answer = (
        "Theo Điểm c Khoản 7 Điều 7 Nghị định 168/2024/NĐ-CP, người điều khiển xe mô tô vượt đèn đỏ "
        "bị phạt từ 4 triệu đến 6 triệu đồng, trừ 4 điểm. Khác với người đi xe ô tô bị phạt từ 18.000.000 đến 20.000.000."
    )
    res = evaluate_generation(
        ground_truth=sample_fine_lookup_item.factual_gt,
        assertions=sample_fine_lookup_item.assertions,
        generated_answer=answer,
    )
    assert res.passed is False
    assert "xe ô tô" in res.forbidden_keywords_detected
    assert "18.000.000" in res.forbidden_keywords_detected


def test_evaluate_generation_regex_unmatched(
    sample_fine_lookup_item: DeterministicBenchmarkItem,
) -> None:
    # Missing points deduction
    answer = (
        "Theo Điểm c Khoản 7 Điều 7 Nghị định 168/2024/NĐ-CP, người điều khiển xe mô tô vượt đèn đỏ "
        "bị phạt từ 4 triệu đến 6 triệu đồng."
    )
    res = evaluate_generation(
        ground_truth=sample_fine_lookup_item.factual_gt,
        assertions=sample_fine_lookup_item.assertions,
        generated_answer=answer,
    )
    assert res.passed is False
    # Regex for points was not matched
    assert res.regex_passed.get(r"trừ\s*(4|bốn)\s*điểm") is False


def test_evaluate_generation_refusal(
    sample_out_of_scope_item: DeterministicBenchmarkItem,
) -> None:
    answer_refusal = (
        "Hệ thống chỉ hỗ trợ tra cứu pháp luật giao thông đường bộ Việt Nam. "
        "Câu hỏi về tàu hỏa đường sắt nằm ngoài phạm vi điều chỉnh của hệ thống."
    )
    res = evaluate_generation(
        ground_truth=sample_out_of_scope_item.factual_gt,
        assertions=sample_out_of_scope_item.assertions,
        generated_answer=answer_refusal,
    )
    assert res.passed is True
    assert res.refusal_passed is True

    answer_non_refusal = "Tàu hỏa vượt đèn đỏ bị phạt 10 triệu đồng."
    res_fail = evaluate_generation(
        ground_truth=sample_out_of_scope_item.factual_gt,
        assertions=sample_out_of_scope_item.assertions,
        generated_answer=answer_non_refusal,
    )
    assert res_fail.passed is False
    assert res_fail.refusal_passed is False


# ============================================================================
# 4. Failure Root-Cause Attribution Tests
# ============================================================================


def test_root_cause_retrieval_forbidden_collision(
    sample_fine_lookup_item: DeterministicBenchmarkItem,
) -> None:
    res = evaluate_benchmark_item(
        item=sample_fine_lookup_item,
        retrieved_unit_ids=["168_2024_ND-CP_D7_K7_Dc", "168_2024_ND-CP_D6_K9_Db"],
        generated_answer="",
        mode="retrieval",
    )
    assert res.overall_passed is False
    assert res.root_cause == FailureRootCause.RETRIEVAL_FORBIDDEN_NODE_COLLISION


def test_root_cause_retrieval_missing_must_have(
    sample_fine_lookup_item: DeterministicBenchmarkItem,
) -> None:
    res = evaluate_benchmark_item(
        item=sample_fine_lookup_item,
        retrieved_unit_ids=["some_unrelated_node"],
        generated_answer="",
        mode="retrieval",
    )
    assert res.overall_passed is False
    assert res.root_cause == FailureRootCause.RETRIEVAL_MISSING_MUST_HAVE


def test_root_cause_generation_forbidden_leak(
    sample_fine_lookup_item: DeterministicBenchmarkItem,
) -> None:
    res = evaluate_benchmark_item(
        item=sample_fine_lookup_item,
        retrieved_unit_ids=["168_2024_ND-CP_D7_K7_Dc"],
        generated_answer="Phạt 4 đến 6 triệu đồng, trừ 4 điểm, nhưng xe ô tô thì phạt nặng hơn.",
        mode="e2e",
    )
    assert res.overall_passed is False
    assert res.root_cause == FailureRootCause.GENERATION_FORBIDDEN_KEYWORD_LEAK


def test_root_cause_generation_sanction_mismatch(
    sample_fine_lookup_item: DeterministicBenchmarkItem,
) -> None:
    res = evaluate_benchmark_item(
        item=sample_fine_lookup_item,
        retrieved_unit_ids=["168_2024_ND-CP_D7_K7_Dc"],
        generated_answer="Phạt từ 800.000 đồng đến 1.000.000 đồng đối với xe mô tô.",
        mode="e2e",
    )
    assert res.overall_passed is False
    assert res.root_cause in (
        FailureRootCause.GENERATION_SANCTION_MISMATCH,
        FailureRootCause.GENERATION_REGEX_UNMATCHED,
    )


def test_root_cause_success_none(
    sample_fine_lookup_item: DeterministicBenchmarkItem,
) -> None:
    perfect_answer = (
        "Theo Điểm c Khoản 7 Điều 7 Nghị định 168/2024/NĐ-CP, người điều khiển xe mô tô "
        "vượt đèn đỏ bị phạt tiền từ 4.000.000 đồng đến 6.000.000 đồng và bị trừ 4 điểm GPLX."
    )
    res = evaluate_benchmark_item(
        item=sample_fine_lookup_item,
        retrieved_unit_ids=["168_2024_ND-CP_D7_K7_Dc"],
        generated_answer=perfect_answer,
        mode="e2e",
    )
    assert res.overall_passed is True
    assert res.root_cause == FailureRootCause.NONE


# ============================================================================
# 5. Orchestrator & Aggregation Tests
# ============================================================================


def test_orchestrator_deterministic_dataset_mocked(
    sample_fine_lookup_item: DeterministicBenchmarkItem,
    sample_out_of_scope_item: DeterministicBenchmarkItem,
) -> None:
    orchestrator = EvaluationOrchestrator()

    # Mock evaluate_deterministic_sample to avoid calling real pipeline
    def mock_eval_sample(item: DeterministicBenchmarkItem, mode: str = "retrieval"):
        if item.test_id == "TC_TEST_001":
            return BenchmarkItemEvaluationResult(
                test_id=item.test_id,
                category=item.category,
                question=item.question,
                mode=mode,
                retrieval_result=RetrievalEvaluationResult(
                    must_have_matched=["168_2024_ND-CP_D7_K7_Dc"],
                    recall=1.0,
                    passed=True,
                ),
                overall_passed=True,
                root_cause=FailureRootCause.NONE,
                execution_time_ms=12.5,
            )
        else:
            return BenchmarkItemEvaluationResult(
                test_id=item.test_id,
                category=item.category,
                question=item.question,
                mode=mode,
                retrieval_result=RetrievalEvaluationResult(
                    must_not_have_detected=["forbidden_node"],
                    collision_rate=1.0,
                    passed=False,
                ),
                overall_passed=False,
                root_cause=FailureRootCause.RETRIEVAL_FORBIDDEN_NODE_COLLISION,
                execution_time_ms=5.0,
            )

    orchestrator.evaluate_deterministic_sample = MagicMock(side_effect=mock_eval_sample)  # type: ignore[method-assign]

    items = [sample_fine_lookup_item, sample_out_of_scope_item]
    summary: BenchmarkRunSummary = orchestrator.evaluate_deterministic_dataset(
        items, mode="retrieval"
    )

    assert summary.total_samples == 2
    assert summary.passed_samples == 1
    assert summary.overall_pass_rate == 0.5
    assert summary.retrieval_pass_rate == 0.5
    assert summary.root_cause_counts["RETRIEVAL_FORBIDDEN_NODE_COLLISION"] == 1
    assert summary.root_cause_counts["NONE"] == 1
    assert "FINE_LOOKUP" in summary.category_breakdown
    assert summary.category_breakdown["FINE_LOOKUP"]["passed"] == 1


# ============================================================================
# 6. Exporter Excel & JSON Tests
# ============================================================================


def test_export_deterministic_benchmark_to_excel(
    tmp_path: Path,
    sample_fine_lookup_item: DeterministicBenchmarkItem,
) -> None:
    res = BenchmarkItemEvaluationResult(
        test_id=sample_fine_lookup_item.test_id,
        category=sample_fine_lookup_item.category,
        question=sample_fine_lookup_item.question,
        mode="retrieval",
        retrieval_result=RetrievalEvaluationResult(
            must_have_matched=["168_2024_ND-CP_D7_K7_Dc"],
            recall=1.0,
            passed=True,
        ),
        overall_passed=True,
        root_cause=FailureRootCause.NONE,
        execution_time_ms=15.0,
    )

    summary = BenchmarkRunSummary(
        timestamp="2026-09-25T12:00:00Z",
        mode="retrieval",
        total_samples=1,
        passed_samples=1,
        overall_pass_rate=1.0,
        retrieval_pass_rate=1.0,
        root_cause_counts={"NONE": 1},
        category_breakdown={"FINE_LOOKUP": {"total": 1, "passed": 1, "pass_rate": 1.0}},
        results=[res],
    )

    excel_file = tmp_path / "test_benchmark_report.xlsx"
    out_path = EvaluationExporter.export_deterministic_benchmark_to_excel(
        summary, excel_file
    )
    assert out_path.exists()

    wb = load_workbook(str(out_path), data_only=True)
    assert "Dashboard Tổng quan" in wb.sheetnames
    assert "Chi tiết Test Cases" in wb.sheetnames

    ws_dash = wb["Dashboard Tổng quan"]
    assert "BÁO CÁO KIỂM CHUẨN TẤT ĐỊNH" in str(ws_dash["A1"].value)

    ws_cases = wb["Chi tiết Test Cases"]
    assert ws_cases["A2"].value == "TC_TEST_001"
    assert ws_cases["D2"].value == "PASS"
