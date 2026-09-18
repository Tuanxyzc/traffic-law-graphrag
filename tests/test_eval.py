"""Unit tests for the RAGAS & Domain-Specific Dual-Evaluation Framework."""

from __future__ import annotations

import json
from pathlib import Path
from unittest.mock import MagicMock

from openpyxl import load_workbook

from src.eval.deterministic import (
    DeterministicEvaluator,
    detect_warning,
    extract_cited_provisions,
    extract_fine_range,
    matches_provision,
    parse_vnd_amount,
)
from src.eval.evaluator import EvaluationOrchestrator
from src.eval.exporter import EvaluationExporter
from src.eval.models import (
    EvaluationCategory,
    EvaluationResult,
    EvaluationSample,
    EvaluationSummary,
    RagasScore,
)
from src.eval.ragas_judge import RagasJudge, RagasJudgeClient, cosine_similarity


# ==========================================
# 1. Tests for Deterministic Evaluator
# ==========================================
def test_parse_vnd_amount() -> None:
    assert parse_vnd_amount("800.000", "đồng") == 800_000
    assert parse_vnd_amount("1.000.000", "đ") == 1_000_000
    assert parse_vnd_amount("18", "triệu") == 18_000_000
    assert parse_vnd_amount("20", "tr") == 20_000_000
    assert parse_vnd_amount("150.000", "đồng") == 150_000
    assert parse_vnd_amount("abc", "đồng") is None


def test_extract_fine_range() -> None:
    text1 = "Người vi phạm bị phạt tiền từ 800.000 đồng đến 1.000.000 đồng."
    assert extract_fine_range(text1) == (800_000, 1_000_000)

    text2 = "Phạt từ 18 triệu đến 20 triệu đồng đối với hành vi vượt đèn đỏ."
    assert extract_fine_range(text2) == (18_000_000, 20_000_000)

    text3 = "Hành vi này bị phạt tiền từ 2 đến 3 triệu đồng."
    assert extract_fine_range(text3) == (2_000_000, 3_000_000)

    text4 = "Áp dụng hình thức phạt tiền 500.000 đồng."
    assert extract_fine_range(text4) == (500_000, 500_000)

    text5 = "Không có chế tài xử phạt tiền đối với hành vi này."
    assert extract_fine_range(text5) == (None, None)


def test_provision_matching() -> None:
    expected = "168_2024_ND-CP_D6_K9_Db"
    candidates = {"168_2024_ND_CP_D6_K9_DB", "168_2024_ND_CP_D7_K7_DC"}
    assert matches_provision(expected, candidates) is True

    # Hierarchy matching (parent clause matches point)
    parent_cand = {"168_2024_ND_CP_D6_K9"}
    assert matches_provision(expected, parent_cand) is True

    non_match = {"168_2024_ND_CP_D7_K1"}
    assert matches_provision(expected, non_match) is False


def test_extract_cited_provisions() -> None:
    citations = ["168_2024_ND-CP_D6_K9_Db"]
    answer = "Căn cứ theo Điều 6 Khoản 9 Điểm b Nghị định 168_2024_ND-CP_D6_K9_Db quy định mức phạt..."
    cited = extract_cited_provisions(citations, answer)
    assert "168_2024_ND_CP_D6_K9_DB" in cited


def test_detect_warning() -> None:
    assert detect_warning("Văn bản này đã hết hiệu lực từ ngày 01/01/2025") is True
    assert detect_warning("Điều khoản này đã được sửa đổi bởi Nghị định 238") is True
    assert detect_warning("Quy định thông thường đang có hiệu lực") is False
    assert detect_warning("Lưu ý hiệu lực từ năm 2026", custom_keywords=["2026"]) is True


def test_deterministic_evaluator_pass() -> None:
    evaluator = DeterministicEvaluator(recall_threshold=0.5)
    sample = EvaluationSample(
        id="TC_TEST_01",
        category=EvaluationCategory.FINE_LOOKUP,
        question="Vượt đèn đỏ ô tô phạt bao nhiêu?",
        expected_provision_ids=["168_2024_ND-CP_D6_K9_Db"],
        expected_fine_min=18_000_000,
        expected_fine_max=20_000_000,
        is_amended_case=False,
        ground_truth_answer="Phạt từ 18 triệu đến 20 triệu theo Điểm b Khoản 9 Điều 6.",
    )

    answer = "Theo 168_2024_ND-CP_D6_K9_Db, phạt tiền từ 18.000.000 đồng đến 20.000.000 đồng."
    citations = ["168_2024_ND-CP_D6_K9_Db"]

    score = evaluator.evaluate(sample, answer, citations)
    assert score.provision_recall == 1.0
    assert score.fine_exact_match is True
    assert score.deterministic_passed is True


def test_deterministic_evaluator_out_of_scope() -> None:
    evaluator = DeterministicEvaluator()
    sample = EvaluationSample(
        id="TC_TEST_OOS",
        category=EvaluationCategory.OUT_OF_SCOPE,
        question="Đi bộ ngõ xóm có bị phạt tiền không?",
        expected_provision_ids=[],
        expected_fine_min=None,
        expected_fine_max=None,
        is_amended_case=False,
        ground_truth_answer="Pháp luật không quy định xử phạt vi phạm hành chính.",
    )

    answer = "Pháp luật hiện hành không quy định xử phạt vi phạm hành chính đối với hành vi này."
    citations: list[str] = []

    score = evaluator.evaluate(sample, answer, citations)
    assert score.provision_recall == 1.0
    assert score.deterministic_passed is True


# ==========================================
# 2. Tests for RAGAS Judge
# ==========================================
def test_cosine_similarity() -> None:
    assert cosine_similarity([1.0, 0.0], [1.0, 0.0]) == 1.0
    assert cosine_similarity([1.0, 0.0], [0.0, 1.0]) == 0.0
    assert cosine_similarity([], [1.0]) == 0.0


def test_ragas_judge_with_mocks() -> None:
    mock_client = MagicMock(spec=RagasJudgeClient)
    mock_embedding = MagicMock()

    # Mock Faithfulness response
    mock_client.generate_json.side_effect = [
        # Call 1: Faithfulness
        {
            "statements": [
                {"statement": "Ô tô vượt đèn đỏ phạt 18-20tr", "supported": True},
                {"statement": "Bị trừ 4 điểm GPLX", "supported": True},
            ]
        },
        # Call 2: Answer Relevancy
        {
            "generated_questions": [
                "Ô tô vượt đèn đỏ phạt bao nhiêu?",
                "Mức phạt ô tô không chấp hành đèn tín hiệu?",
                "Vượt đèn đỏ ô tô bị trừ mấy điểm?",
            ]
        },
        # Call 3: Context Precision
        {
            "contexts_relevance": [
                {"index": 1, "is_relevant": True},
                {"index": 2, "is_relevant": False},
            ]
        },
        # Call 4: Context Recall
        {
            "ground_truth_statements": [
                {"statement": "Phạt tiền từ 18 đến 20 triệu", "attributed": True},
                {"statement": "Quy định tại Nghị định 168", "attributed": True},
            ]
        },
    ]

    mock_embedding.embed_query.return_value = [1.0, 0.0]
    mock_embedding.embed_texts.return_value = [[1.0, 0.0], [0.9, 0.1], [0.8, 0.2]]

    judge = RagasJudge(
        judge_client=mock_client,
        embedding_manager=mock_embedding,
        faithfulness_threshold=0.7,
        relevancy_threshold=0.6,
        context_precision_threshold=0.6,
        context_recall_threshold=0.6,
    )

    sample = EvaluationSample(
        id="TC_MOCK",
        category=EvaluationCategory.FINE_LOOKUP,
        question="Ô tô vượt đèn đỏ phạt bao nhiêu?",
        ground_truth_answer="Phạt từ 18 đến 20 triệu theo Nghị định 168.",
    )

    score = judge.evaluate(
        sample=sample,
        generated_answer="Ô tô vượt đèn đỏ phạt 18-20tr và trừ 4 điểm.",
        retrieved_contexts=["Điều 6 quy định phạt 18-20tr.", "Điều 7 quy định xe máy."],
    )

    assert score.faithfulness == 1.0
    assert score.answer_relevancy is not None and score.answer_relevancy > 0.8
    assert score.context_precision == 1.0
    assert score.context_recall == 1.0
    assert score.ragas_passed is True


# ==========================================
# 3. Tests for Exporter & Dataset Tooling
# ==========================================
def test_exporter_excel_and_csv(tmp_path: Path) -> None:
    sample = EvaluationSample(
        id="TC_EXP",
        category=EvaluationCategory.FINE_LOOKUP,
        question="Xe máy vượt đèn đỏ phạt bao nhiêu?",
        expected_provision_ids=["168_2024_ND-CP_D7_K7_Dc"],
        expected_fine_min=4_000_000,
        expected_fine_max=6_000_000,
        ground_truth_answer="Phạt từ 4 đến 6 triệu.",
    )

    det_score = DeterministicEvaluator().evaluate(
        sample=sample,
        generated_answer="Theo 168_2024_ND-CP_D7_K7_Dc, phạt tiền từ 4.000.000 đồng đến 6.000.000 đồng.",
        citations=["168_2024_ND-CP_D7_K7_Dc"],
    )

    res = EvaluationResult(
        sample=sample,
        generated_answer="Phạt từ 4.000.000 đồng đến 6.000.000 đồng.",
        citations=["168_2024_ND-CP_D7_K7_Dc"],
        retrieved_contexts=["Trích đoạn Điều 7 Khoản 7 Điểm c."],
        retrieved_unit_ids=["168_2024_ND-CP_D7_K7_Dc"],
        deterministic_score=det_score,
        ragas_score=RagasScore(
            faithfulness=1.0,
            answer_relevancy=0.95,
            context_precision=1.0,
            context_recall=1.0,
            ragas_passed=True,
        ),
        latency_ms=150.0,
        overall_passed=True,
    )

    summary = EvaluationSummary(
        timestamp="2026-09-15T18:00:00",
        tier="full",
        total_samples=1,
        passed_samples=1,
        overall_pass_rate=1.0,
        avg_provision_recall=1.0,
        avg_provision_precision=1.0,
        avg_provision_f1=1.0,
        fine_accuracy=1.0,
        warning_accuracy=1.0,
        avg_faithfulness=1.0,
        avg_answer_relevancy=0.95,
        avg_context_precision=1.0,
        avg_context_recall=1.0,
        avg_latency_ms=150.0,
        results=[res],
    )

    # 1. Test Excel export
    excel_path = tmp_path / "test_report.xlsx"
    EvaluationExporter.export_to_excel(summary, excel_path)
    assert excel_path.exists()

    wb = load_workbook(excel_path)
    assert "Dashboard Tổng quan" in wb.sheetnames
    assert "Chi tiết từng câu hỏi" in wb.sheetnames
    ws_det = wb["Chi tiết từng câu hỏi"]
    assert ws_det.cell(row=2, column=1).value == "TC_EXP"
    assert ws_det.cell(row=2, column=4).value == "PASS"

    # 2. Test CSV export
    csv_path = tmp_path / "test_report.csv"
    EvaluationExporter.export_to_csv(summary, csv_path)
    assert csv_path.exists()
    content = csv_path.read_text(encoding="utf-8-sig")
    assert "TC_EXP" in content
    assert "PASS" in content


def test_dataset_excel_roundtrip(tmp_path: Path) -> None:
    sample = EvaluationSample(
        id="TC_ROUNDTRIP",
        category=EvaluationCategory.TEMPORAL_VALIDITY,
        question="Nghị định 168 có hiệu lực khi nào?",
        expected_provision_ids=["168_2024_ND-CP_D53_K1"],
        expected_fine_min=None,
        expected_fine_max=None,
        expected_points_deducted=None,
        is_amended_case=False,
        expected_warning_keywords=["01/01/2025"],
        ground_truth_answer="Hiệu lực từ ngày 01/01/2025.",
        notes="Ghi chú kiểm thử",
    )

    excel_file = tmp_path / "dataset_roundtrip.xlsx"
    EvaluationExporter.export_dataset_to_excel([sample], excel_file)
    assert excel_file.exists()

    imported = EvaluationExporter.import_dataset_from_excel(excel_file)
    assert len(imported) == 1
    assert imported[0].id == "TC_ROUNDTRIP"
    assert imported[0].category == EvaluationCategory.TEMPORAL_VALIDITY
    assert imported[0].expected_provision_ids == ["168_2024_ND-CP_D53_K1"]
    assert imported[0].expected_warning_keywords == ["01/01/2025"]
    assert imported[0].ground_truth_answer == "Hiệu lực từ ngày 01/01/2025."


# ==========================================
# 4. Tests for Evaluation Orchestrator
# ==========================================
def test_evaluation_orchestrator_mocked() -> None:
    orchestrator = EvaluationOrchestrator(tier="deterministic")
    sample = EvaluationSample(
        id="TC_ORCH",
        category=EvaluationCategory.FINE_LOOKUP,
        question="Xe đạp vượt đèn đỏ phạt bao nhiêu?",
        expected_provision_ids=["168_2024_ND-CP_D9_K2_Dd"],
        expected_fine_min=150_000,
        expected_fine_max=250_000,
        ground_truth_answer="Phạt từ 150.000 đến 250.000 đồng.",
    )

    result = orchestrator.evaluate_sample(
        sample=sample,
        mock_answer="Theo Điểm đ Khoản 2 Điều 9, phạt tiền từ 150.000 đồng đến 250.000 đồng.",
        mock_citations=["168_2024_ND-CP_D9_K2_Dd"],
        mock_contexts=["Quy định mức phạt xe đạp..."],
        mock_unit_ids=["168_2024_ND-CP_D9_K2_Dd"],
    )

    assert result.overall_passed is True
    assert result.deterministic_score.provision_recall == 1.0
    assert result.deterministic_score.fine_exact_match is True


def test_golden_dataset_validation() -> None:
    golden_path = Path("data/eval/golden_dataset.json")
    assert golden_path.exists()

    with open(golden_path, "r", encoding="utf-8") as f:
        raw_data = json.load(f)

    assert len(raw_data) >= 25
    samples = [EvaluationSample.model_validate(item) for item in raw_data]

    categories = {s.category for s in samples}
    assert EvaluationCategory.FINE_LOOKUP in categories
    assert EvaluationCategory.TEMPORAL_VALIDITY in categories
    assert EvaluationCategory.DOCUMENT_AMENDMENT in categories
    assert EvaluationCategory.MULTI_HOP_RULE_SANCTION in categories
    assert EvaluationCategory.OUT_OF_SCOPE in categories
