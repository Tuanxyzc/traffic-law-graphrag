"""Reporting and Exporter module for Evaluation results: Excel (openpyxl) and CSV."""

from __future__ import annotations

import csv
import logging
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.eval.models import (
    EvaluationCategory,
    EvaluationSample,
    EvaluationSummary,
)

logger = logging.getLogger(__name__)

# Styling Constants for Excel
NAVY_HEADER_FILL = PatternFill(
    start_color="1F4E79", end_color="1F4E79", fill_type="solid"
)
BLUE_SUBHEADER_FILL = PatternFill(
    start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"
)
PASS_FILL = PatternFill(
    start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"
)  # Light green
FAIL_FILL = PatternFill(
    start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"
)  # Light red/orange
WHITE_BOLD_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
DARK_BOLD_FONT = Font(name="Calibri", size=11, bold=True, color="000000")
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="1F4E79")
REGULAR_FONT = Font(name="Calibri", size=10)
PASS_FONT = Font(name="Calibri", size=10, bold=True, color="375623")
FAIL_FONT = Font(name="Calibri", size=10, bold=True, color="C65911")

THIN_BORDER_SIDE = Side(border_style="thin", color="D9D9D9")
GRID_BORDER = Border(
    left=THIN_BORDER_SIDE,
    right=THIN_BORDER_SIDE,
    top=THIN_BORDER_SIDE,
    bottom=THIN_BORDER_SIDE,
)


def format_vnd(amount: int | None) -> str:
    """Formats integer VND to readable text (e.g. 800.000đ)."""
    if amount is None:
        return "-"
    return f"{amount:,}đ".replace(",", ".")


def auto_fit_columns(sheet: Any, max_width: int = 50) -> None:
    """Adjusts column widths based on content length."""
    for col in sheet.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            first_line = val.split("\n")[0]
            max_len = max(max_len, len(first_line))
        adjusted_width = min(max(max_len + 3, 10), max_width)
        sheet.column_dimensions[col_letter].width = adjusted_width


class EvaluationExporter:
    """Generates formatted reports in Excel and CSV from evaluation runs."""

    @staticmethod
    def export_to_excel(summary: EvaluationSummary, output_path: str | Path) -> Path:
        """Exports evaluation summary and detailed results into a styled Excel workbook."""
        out_file = Path(output_path)
        out_file.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()

        # ==========================================
        # Sheet 1: Dashboard Summary
        # ==========================================
        ws_dash = wb.active
        ws_dash.title = "Dashboard Tổng quan"
        ws_dash.views.sheetView[0].showGridLines = True

        # 1. Title Block
        ws_dash.merge_cells("A1:F1")
        ws_dash["A1"] = "BÁO CÁO ĐÁNH GIÁ CHẤT LƯỢNG HỆ THỐNG GRAPHRAG LUẬT GIAO THÔNG"
        ws_dash["A1"].font = TITLE_FONT
        ws_dash["A1"].alignment = Alignment(vertical="center")
        ws_dash.row_dimensions[1].height = 35

        ws_dash["A2"] = f"Thời gian đánh giá: {summary.timestamp}"
        ws_dash["A2"].font = REGULAR_FONT
        ws_dash["A3"] = f"Chế độ đánh giá: {summary.tier.upper()}"
        ws_dash["A3"].font = REGULAR_FONT
        ws_dash["A4"] = f"Tổng số câu hỏi kiểm thử: {summary.total_samples}"
        ws_dash["A4"].font = REGULAR_FONT

        # 2. Key Metrics Table
        headers_kpi = [
            "Chỉ số đánh giá",
            "Giá trị",
            "Mục tiêu tối thiểu",
            "Mô tả ý nghĩa",
        ]
        row_kpi_start = 6
        for col_idx, h in enumerate(headers_kpi, start=1):
            cell = ws_dash.cell(row=row_kpi_start, column=col_idx, value=h)
            cell.font = WHITE_BOLD_FONT
            cell.fill = NAVY_HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = GRID_BORDER
        ws_dash.row_dimensions[row_kpi_start].height = 24

        kpi_rows = [
            (
                "Tỉ lệ Pass toàn diện (Overall Pass Rate)",
                f"{summary.overall_pass_rate * 100:.1f}%",
                ">= 80%",
                "Tỉ lệ câu trả lời thỏa mãn cả tiêu chuẩn luật và chất lượng RAGAS",
            ),
            (
                "Độ bao phủ điều luật (Provision Recall@K)",
                f"{summary.avg_provision_recall * 100:.1f}%",
                ">= 70%",
                "Tỉ lệ các điều/khoản luật kỳ vọng được hệ thống trích dẫn thành công",
            ),
            (
                "Độ chính xác điều luật (Provision Precision@K)",
                f"{summary.avg_provision_precision * 100:.1f}%",
                ">= 60%",
                "Tỉ lệ các điều khoản trích dẫn là chính xác, không trích dẫn thừa",
            ),
            (
                "Độ chính xác mức phạt tiền (Fine Range Accuracy)",
                f"{summary.fine_accuracy * 100:.1f}%",
                ">= 85%",
                "Tỉ lệ trích xuất đúng 100% khung tiền phạt tối thiểu và tối đa",
            ),
            (
                "Nhận diện cảnh báo hiệu lực/sửa đổi (Warning Accuracy)",
                f"{summary.warning_accuracy * 100:.1f}%",
                ">= 90%",
                "Tỉ lệ cảnh báo đúng các văn bản đã bị sửa đổi, thay thế hoặc hết hiệu lực",
            ),
            (
                "Điểm trung thực căn cứ (Faithfulness)",
                f"{summary.avg_faithfulness * 100:.1f}%"
                if summary.avg_faithfulness is not None
                else "N/A",
                ">= 70%",
                "Các nhận định trong câu trả lời phải được bảo chứng bởi ngữ cảnh truy xuất",
            ),
            (
                "Độ tương quan câu trả lời (Answer Relevancy)",
                f"{summary.avg_answer_relevancy * 100:.1f}%"
                if summary.avg_answer_relevancy is not None
                else "N/A",
                ">= 65%",
                "Câu trả lời giải đáp đúng trọng tâm câu hỏi của người dân",
            ),
            (
                "Độ chính xác ngữ cảnh (Context Precision)",
                f"{summary.avg_context_precision * 100:.1f}%"
                if summary.avg_context_precision is not None
                else "N/A",
                ">= 60%",
                "Các đoạn ngữ cảnh liên quan được xếp hạng ưu tiên ở vị trí cao",
            ),
            (
                "Độ bao phủ ngữ cảnh (Context Recall)",
                f"{summary.avg_context_recall * 100:.1f}%"
                if summary.avg_context_recall is not None
                else "N/A",
                ">= 60%",
                "Ngữ cảnh truy xuất chứa đầy đủ thông tin để trả lời câu hỏi chuẩn",
            ),
            (
                "Thời gian phản hồi trung bình (Average Latency)",
                f"{summary.avg_latency_ms:.1f} ms",
                "< 4000 ms",
                "Thời gian hoàn thành xử lý toàn diện qua GraphRAG pipeline",
            ),
        ]

        for offset, (name, val, target, desc) in enumerate(kpi_rows, start=1):
            r = row_kpi_start + offset
            c1 = ws_dash.cell(row=r, column=1, value=name)
            c2 = ws_dash.cell(row=r, column=2, value=val)
            c3 = ws_dash.cell(row=r, column=3, value=target)
            c4 = ws_dash.cell(row=r, column=4, value=desc)
            for c in (c1, c2, c3, c4):
                c.font = REGULAR_FONT
                c.border = GRID_BORDER
            c1.font = DARK_BOLD_FONT
            c2.alignment = Alignment(horizontal="center")
            c3.alignment = Alignment(horizontal="center")
            ws_dash.row_dimensions[r].height = 20

        # 3. Category Breakdown Table
        row_cat_start = row_kpi_start + len(kpi_rows) + 3
        ws_dash.cell(
            row=row_cat_start - 1,
            column=1,
            value="PHÂN TÍCH HIỆU NĂNG THEO NHÓM CÂU HỎI",
        ).font = Font(name="Calibri", size=13, bold=True, color="1F4E79")

        cat_headers = [
            "Nhóm câu hỏi",
            "Số mẫu",
            "Số mẫu đạt",
            "Tỉ lệ Đạt",
            "Recall Điều luật",
            "Precision Điều luật",
            "Faithfulness",
            "Relevancy",
        ]
        for col_idx, h in enumerate(cat_headers, start=1):
            cell = ws_dash.cell(row=row_cat_start, column=col_idx, value=h)
            cell.font = WHITE_BOLD_FONT
            cell.fill = NAVY_HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = GRID_BORDER

        cat_offset = 1
        for cat_name, metrics in summary.category_breakdown.items():
            r = row_cat_start + cat_offset
            ws_dash.cell(row=r, column=1, value=cat_name).font = DARK_BOLD_FONT
            ws_dash.cell(
                row=r, column=2, value=metrics.total_samples
            ).alignment = Alignment(horizontal="center")
            ws_dash.cell(
                row=r, column=3, value=metrics.passed_samples
            ).alignment = Alignment(horizontal="center")
            ws_dash.cell(
                row=r, column=4, value=f"{metrics.pass_rate * 100:.1f}%"
            ).alignment = Alignment(horizontal="center")
            ws_dash.cell(
                row=r, column=5, value=f"{metrics.avg_provision_recall * 100:.1f}%"
            ).alignment = Alignment(horizontal="center")
            ws_dash.cell(
                row=r, column=6, value=f"{metrics.avg_provision_precision * 100:.1f}%"
            ).alignment = Alignment(horizontal="center")
            ws_dash.cell(
                row=r,
                column=7,
                value=f"{metrics.avg_faithfulness * 100:.1f}%"
                if metrics.avg_faithfulness is not None
                else "-",
            ).alignment = Alignment(horizontal="center")
            ws_dash.cell(
                row=r,
                column=8,
                value=f"{metrics.avg_answer_relevancy * 100:.1f}%"
                if metrics.avg_answer_relevancy is not None
                else "-",
            ).alignment = Alignment(horizontal="center")
            for col in range(1, 9):
                ws_dash.cell(row=r, column=col).border = GRID_BORDER
            cat_offset += 1

        auto_fit_columns(ws_dash, max_width=60)

        # ==========================================
        # Sheet 2: Detailed Results
        # ==========================================
        ws_det = wb.create_sheet(title="Chi tiết từng câu hỏi")
        ws_det.views.sheetView[0].showGridLines = True

        det_headers = [
            "ID",
            "Danh mục",
            "Câu hỏi của công dân",
            "Kết quả",
            "Điều luật kỳ vọng",
            "Điều luật trích dẫn",
            "Recall Điều luật",
            "Khung phạt kỳ vọng",
            "Khung phạt trả về",
            "Đúng mức phạt?",
            "Cảnh báo hiệu lực",
            "Faithfulness",
            "Answer Relevancy",
            "Context Precision",
            "Context Recall",
            "Độ trễ (ms)",
            "Câu trả lời bot",
            "Đáp án chuẩn (Ground Truth)",
            "Ngữ cảnh trích xuất",
        ]

        for col_idx, h in enumerate(det_headers, start=1):
            cell = ws_det.cell(row=1, column=col_idx, value=h)
            cell.font = WHITE_BOLD_FONT
            cell.fill = NAVY_HEADER_FILL
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.border = GRID_BORDER
        ws_det.row_dimensions[1].height = 28

        for row_idx, res in enumerate(summary.results, start=2):
            sample = res.sample
            det_score = res.deterministic_score
            rag_score = res.ragas_score

            # Fine display
            exp_fine_str = (
                f"{format_vnd(sample.expected_fine_min)} - {format_vnd(sample.expected_fine_max)}"
                if sample.expected_fine_min or sample.expected_fine_max
                else "-"
            )
            act_fine_str = (
                f"{format_vnd(det_score.extracted_fine_min)} - {format_vnd(det_score.extracted_fine_max)}"
                if det_score.extracted_fine_min or det_score.extracted_fine_max
                else "-"
            )

            fine_match_str = (
                "ĐẠT"
                if det_score.fine_exact_match is True
                else ("SAI" if det_score.fine_exact_match is False else "N/A")
            )

            status_str = "PASS" if res.overall_passed else "FAIL"

            row_values = [
                sample.id,
                sample.category.value,
                sample.question,
                status_str,
                ", ".join(sample.expected_provision_ids) or "Không",
                ", ".join(res.citations) or "Không",
                f"{det_score.provision_recall * 100:.1f}%",
                exp_fine_str,
                act_fine_str,
                fine_match_str,
                "ĐẠT" if det_score.warning_match else "CHƯA",
                f"{rag_score.faithfulness * 100:.1f}%"
                if rag_score and rag_score.faithfulness is not None
                else "-",
                f"{rag_score.answer_relevancy * 100:.1f}%"
                if rag_score and rag_score.answer_relevancy is not None
                else "-",
                f"{rag_score.context_precision * 100:.1f}%"
                if rag_score and rag_score.context_precision is not None
                else "-",
                f"{rag_score.context_recall * 100:.1f}%"
                if rag_score and rag_score.context_recall is not None
                else "-",
                f"{res.latency_ms:.1f}",
                res.generated_answer,
                sample.ground_truth_answer,
                "\n---\n".join(res.retrieved_contexts[:3]),
            ]

            for col_idx, val in enumerate(row_values, start=1):
                cell = ws_det.cell(row=row_idx, column=col_idx, value=val)
                cell.font = REGULAR_FONT
                cell.border = GRID_BORDER
                cell.alignment = Alignment(vertical="top", wrap_text=True)

            # Color highlight for overall status
            status_cell = ws_det.cell(row=row_idx, column=4)
            if res.overall_passed:
                status_cell.fill = PASS_FILL
                status_cell.font = PASS_FONT
                status_cell.alignment = Alignment(
                    horizontal="center", vertical="center"
                )
            else:
                status_cell.fill = FAIL_FILL
                status_cell.font = FAIL_FONT
                status_cell.alignment = Alignment(
                    horizontal="center", vertical="center"
                )

            ws_det.row_dimensions[row_idx].height = 40

        ws_det.freeze_panes = "A2"
        ws_det.auto_filter.ref = ws_det.dimensions
        auto_fit_columns(ws_det, max_width=45)

        wb.save(str(out_file))
        logger.info("Evaluation report successfully written to Excel: %s", out_file)
        return out_file

    @staticmethod
    def export_to_csv(summary: EvaluationSummary, output_path: str | Path) -> Path:
        """Exports evaluation summary into flat CSV for CI/CD tracking."""
        out_file = Path(output_path)
        out_file.parent.mkdir(parents=True, exist_ok=True)

        headers = [
            "id",
            "category",
            "question",
            "overall_passed",
            "provision_recall",
            "provision_precision",
            "fine_exact_match",
            "warning_match",
            "faithfulness",
            "answer_relevancy",
            "context_precision",
            "context_recall",
            "latency_ms",
            "citations",
            "expected_provisions",
        ]

        with open(out_file, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(headers)

            for res in summary.results:
                det = res.deterministic_score
                rag = res.ragas_score
                writer.writerow(
                    [
                        res.sample.id,
                        res.sample.category.value,
                        res.sample.question,
                        "PASS" if res.overall_passed else "FAIL",
                        det.provision_recall,
                        det.provision_precision,
                        det.fine_exact_match,
                        det.warning_match,
                        rag.faithfulness if rag else "",
                        rag.answer_relevancy if rag else "",
                        rag.context_precision if rag else "",
                        rag.context_recall if rag else "",
                        res.latency_ms,
                        ";".join(res.citations),
                        ";".join(res.sample.expected_provision_ids),
                    ]
                )

        logger.info("Evaluation report successfully written to CSV: %s", out_file)
        return out_file

    @staticmethod
    def export_dataset_to_excel(
        samples: list[EvaluationSample], output_path: str | Path
    ) -> Path:
        """Exports evaluation dataset to Excel for easy editing by human evaluators."""
        out_file = Path(output_path)
        out_file.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = "Golden Dataset"
        ws.views.sheetView[0].showGridLines = True

        headers = [
            "ID",
            "Category",
            "Question",
            "Expected Provisions (semicolon-separated)",
            "Expected Fine Min (VND)",
            "Expected Fine Max (VND)",
            "Expected Points Deducted",
            "Is Amended Case (TRUE/FALSE)",
            "Expected Warning Keywords (semicolon-separated)",
            "Ground Truth Answer",
            "Notes",
        ]

        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.font = WHITE_BOLD_FONT
            cell.fill = NAVY_HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = GRID_BORDER
        ws.row_dimensions[1].height = 25

        for row_idx, s in enumerate(samples, start=2):
            row_vals = [
                s.id,
                s.category.value,
                s.question,
                ";".join(s.expected_provision_ids),
                s.expected_fine_min,
                s.expected_fine_max,
                s.expected_points_deducted,
                s.is_amended_case,
                ";".join(s.expected_warning_keywords),
                s.ground_truth_answer,
                s.notes or "",
            ]
            for col_idx, val in enumerate(row_vals, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = REGULAR_FONT
                cell.border = GRID_BORDER
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            ws.row_dimensions[row_idx].height = 30

        auto_fit_columns(ws, max_width=45)
        wb.save(str(out_file))
        logger.info(
            "Exported %d evaluation samples to Excel: %s", len(samples), out_file
        )
        return out_file

    @staticmethod
    def import_dataset_from_excel(input_path: str | Path) -> list[EvaluationSample]:
        """Imports evaluation samples from Excel workbook, validating each against Pydantic schema."""
        in_file = Path(input_path)
        if not in_file.exists():
            raise FileNotFoundError(f"Dataset file not found: {in_file}")

        wb = load_workbook(str(in_file), data_only=True)
        ws = wb.active
        samples: list[EvaluationSample] = []

        # Read row by row starting from row 2
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue

            tc_id = str(row[0]).strip()
            category_raw = str(row[1]).strip().upper()
            question = str(row[2]).strip()
            exp_provs = [p.strip() for p in str(row[3] or "").split(";") if p.strip()]

            fine_min = int(row[4]) if row[4] not in (None, "") else None
            fine_max = int(row[5]) if row[5] not in (None, "") else None
            points = int(row[6]) if row[6] not in (None, "") else None

            is_amended_val = row[7]
            is_amended = bool(
                is_amended_val is True or str(is_amended_val).strip().upper() == "TRUE"
            )

            keywords = [k.strip() for k in str(row[8] or "").split(";") if k.strip()]
            ground_truth = str(row[9] or "").strip()
            notes = str(row[10]).strip() if len(row) > 10 and row[10] else None

            sample = EvaluationSample(
                id=tc_id,
                category=EvaluationCategory(category_raw),
                question=question,
                expected_provision_ids=exp_provs,
                expected_fine_min=fine_min,
                expected_fine_max=fine_max,
                expected_points_deducted=points,
                is_amended_case=is_amended,
                expected_warning_keywords=keywords,
                ground_truth_answer=ground_truth,
                notes=notes,
            )
            samples.append(sample)

        logger.info(
            "Successfully imported %d evaluation samples from %s", len(samples), in_file
        )
        return samples
